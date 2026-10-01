from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR

OUT = "Deepfakes_e_desinformacao_Maratona_Tech.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

NAVY = RGBColor(12, 24, 48)
PANEL = RGBColor(22, 41, 72)
CYAN = RGBColor(54, 218, 224)
CORAL = RGBColor(255, 105, 96)
WHITE = RGBColor(244, 248, 255)
MUTED = RGBColor(173, 190, 215)
GREEN = RGBColor(90, 224, 170)
FONT = "Aptos"

def rect(slide, x,y,w,h, color, radius=False, line=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
          Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=color
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb=line
    if radius:
        shape.adjustments[0]=0.12
    return shape

def text(slide, value, x,y,w,h, size=18, color=WHITE, bold=False, font=FONT, align=None, valign=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Pt(0); tf.margin_right=Pt(0); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0); tf.vertical_anchor=valign
    for i,line in enumerate(value.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
        if align is not None: p.alignment=align
        p.space_after=Pt(5)
    return box

def base(kicker, title, n, subtitle=None):
    s=prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
    rect(s,0,0,0.13,7.5,CYAN)
    text(s,kicker.upper(),0.65,0.37,10,0.28,11,CYAN,True)
    text(s,title,0.65,0.82,11.9,0.8,30,WHITE,True)
    if subtitle: text(s,subtitle,0.67,1.63,11.7,0.55,14,MUTED)
    rect(s,0.65,7.02,12.05,0.012,PANEL)
    text(s,"MARATONA TECH 2026  •  3º D",0.67,7.12,8,0.2,9,MUTED,True)
    text(s,f"{n:02d}",12.08,7.1,0.55,0.22,10,CYAN,True,align=PP_ALIGN.RIGHT)
    return s

def card(s,x,y,w,h,head,body,accent=CYAN,body_size=17):
    rect(s,x,y,w,h,PANEL,True)
    rect(s,x,y,0.07,h,accent)
    text(s,head,x+0.24,y+0.2,w-0.48,0.38,17,accent,True)
    text(s,body,x+0.24,y+0.73,w-0.48,h-0.92,body_size,WHITE)

# 1 Cover
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
rect(s,0,0,0.15,7.5,CYAN)
for x,y,d,c in [(9.6,0.7,2.2,PANEL),(10.8,2.5,1.3,CYAN),(8.6,3.4,1.1,CORAL),(11.1,4.4,2.0,PANEL),(9.4,5.8,.7,GREEN)]:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background()
text(s,"MARATONA TECH 2026  •  FASE 1",0.8,0.65,7.9,0.4,13,CYAN,True)
text(s,"DEEPFAKES\nE DESINFORMAÇÃO",0.8,1.35,8.5,2.2,36,WHITE,True)
text(s,"Quando a inteligência artificial pode\nmanipular o que parece real",0.83,3.75,7.7,0.95,21,MUTED)
rect(s,0.82,5.2,7.9,0.72,PANEL,True)
text(s,"PARE  •  VERIFIQUE  •  PENSE  •  COMPARTILHE",1.05,5.42,7.45,0.3,14,CYAN,True)
text(s,"Colégio Estadual Marechal Costa e Silva  |  Cidade Gaúcha – PR  |  3º ano D",0.83,6.45,9,0.34,12,WHITE,True)
text(s,"Equipe: Eduardo Valentin Carneiro • Luiz Gabriel Alves dos Santos • Thiago Henrique Malta Garcia\nFelype Justino da Silva • Paulo Ricardo Nunes da Silva • Adryan Vinicius de Almeida",0.83,6.83,11.5,0.52,9,MUTED)

#2 problem
s=base("01 • O problema","Um vídeo pode parecer prova. Mas será?",2,"Deepfakes são conteúdos de imagem, vídeo ou áudio alterados ou gerados por inteligência artificial.")
card(s,.7,2.45,3.8,3.55,"O que está em jogo","Uma pessoa pode parecer dizer ou fazer algo que nunca aconteceu.",CYAN,19)
card(s,4.75,2.45,3.8,3.55,"Por que importa","O conteúdo pode enganar, prejudicar reputações e influenciar decisões.",CORAL,19)
card(s,8.8,2.45,3.8,3.55,"O desafio","Reconhecer sinais, confirmar a origem e reduzir o compartilhamento impulsivo.",GREEN,19)

#3 chain
s=base("02 • Como circula","Da criação ao compartilhamento",3,"A manipulação pode ganhar alcance quando é publicada sem contexto e replicada por muitas pessoas.")
steps=[("1","DADOS E\nFERRAMENTAS","Imagens, voz e IA"),("2","CONTEÚDO\nALTERADO","Corte, edição ou síntese"),("3","PUBLICAÇÃO","Legenda provoca reação"),("4","RECOMENDAÇÃO","Interações ampliam alcance"),("5","COMPARTILHAMENTO","Pessoas repassam")]
for i,(num,head,sub) in enumerate(steps):
    x=.68+i*2.53
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+.68), Inches(2.55), Inches(.64), Inches(.64)); sh.fill.solid(); sh.fill.fore_color.rgb=CYAN if i<3 else CORAL; sh.line.fill.background()
    text(s,num,x+.68,2.72,.64,.24,15,NAVY,True,align=PP_ALIGN.CENTER)
    text(s,head,x,3.42,2.05,.62,15,WHITE,True,align=PP_ALIGN.CENTER)
    text(s,sub,x,4.25,2.05,.78,13,MUTED,False,align=PP_ALIGN.CENTER)
    if i<4: text(s,"→",x+1.95,2.68,.5,.35,22,CYAN,True,align=PP_ALIGN.CENTER)
rect(s,.8,5.55,11.7,.65,PANEL,True)
text(s,"Algoritmos podem recomendar conteúdos com base em sinais de interesse; isso pode ajudar a ampliar a circulação, mas não prova que um vídeo seja verdadeiro.",1.05,5.75,11.2,.34,14,WHITE)

#4 artifact analysis
s=base("03 • Análise do artefato","Estudo de caso: um vídeo sem contexto",4,"Exemplo hipotético para análise — não representa uma pessoa ou acontecimento real.")
rect(s,.72,2.45,5.25,3.72,PANEL,True)
text(s,"“URGENTE: autoridade anuncia\nmedida chocante”",1.05,2.9,4.55,1.1,23,CORAL,True)
text(s,"Vídeo curto, sem data ou link da gravação completa. O áudio parece estranho e a legenda pede compartilhamento imediato.",1.05,4.38,4.5,1.2,17,WHITE)
card(s,6.25,2.45,5.95,1.12,"Sinal de alerta","Urgência + ausência de fonte verificável.",CORAL,16)
card(s,6.25,3.78,5.95,1.12,"O que ainda não sabemos","Quem publicou? Quando? Há versão completa?",CYAN,16)
card(s,6.25,5.11,5.95,1.12,"Conclusão responsável","Indício não é prova. Investigue antes de afirmar.",GREEN,16)

#5 data and risks
s=base("04 • Dados e impactos","Por que deepfakes podem convencer?",5,"Imagens, vídeos e gravações públicas podem ser usados para imitar aparência e voz.")
card(s,.72,2.45,3.8,3.55,"Rastro digital","Fotos, vídeos e áudios publicados ajudam a compor a presença de uma pessoa na rede.",CYAN,17)
card(s,4.77,2.45,3.8,3.55,"Uso indevido","A semelhança pode ser usada sem consentimento para enganar ou constranger.",CORAL,17)
card(s,8.82,2.45,3.8,3.55,"Consequências","Danos à reputação, golpes, assédio e desconfiança sobre conteúdos verdadeiros.",GREEN,17)

#6 root cause
s=base("05 • Causa do problema","A tecnologia é só uma parte da história",6,"A desinformação se fortalece quando conteúdo convincente encontra pressa, pouca verificação e ampla circulação.")
causes=[("FERRAMENTAS","IA facilita criar ou alterar conteúdos"),("INTENÇÃO","Alguém usa o material para enganar"),("CONTEXTO","Legenda omite origem ou data"),("DINÂMICA","Reações e repasses aceleram a difusão")]
for i,(h,b) in enumerate(causes):
    x=.75+(i%2)*6.15; y=2.5+(i//2)*1.75
    card(s,x,y,5.7,1.35,h,b,[CYAN,CORAL,GREEN,CYAN][i],16)
text(s,"Resposta: combinar educação midiática, proteção de dados e responsabilidade de quem cria e compartilha.",1.0,6.25,11.4,.43,17,WHITE,True,align=PP_ALIGN.CENTER)

#7 investigation method
s=base("06 • Como analisar","Quatro perguntas antes de concluir",7,"Uma checagem simples ajuda a separar suspeita de evidência.")
for i,(n,h,b) in enumerate([
    ("01","ORIGEM","Quem publicou? Existe fonte original?"),
    ("02","CONTEXTO","Quando e onde foi gravado? Há versão completa?"),
    ("03","CONFIRMAÇÃO","Outras fontes confiáveis confirmam?"),
    ("04","EVIDÊNCIA","Há sinais consistentes ou só impressão?")]):
    x=.75+(i%2)*6.1; y=2.4+(i//2)*1.72
    rect(s,x,y,5.72,1.35,PANEL,True)
    text(s,n,x+.22,y+.24,.65,.45,23,CYAN,True)
    text(s,h,x+1.0,y+.22,4.4,.3,15,WHITE,True)
    text(s,b,x+1.0,y+.67,4.4,.5,14,MUTED)
rect(s,1.1,6.15,11.1,.57,CORAL,True)
text(s,"Não conclua que é falso só porque “parece estranho”. Verifique a procedência.",1.35,6.32,10.6,.25,14,NAVY,True,align=PP_ALIGN.CENTER)

#8 guide
s=base("07 • Recomendações públicas","PARE. VERIFIQUE. PENSE. COMPARTILHE.",8,"Um guia rápido para agir com responsabilidade diante de um vídeo suspeito.")
items=[("PARE","Não repasse no impulso.",CORAL),("VERIFIQUE","Busque a fonte original e a data.",CYAN),("PENSE","Compare com fontes confiáveis.",GREEN),("DECIDA","Compartilhe só com contexto confirmado.",CYAN)]
for i,(h,b,c) in enumerate(items):
    x=.72+i*3.15
    rect(s,x,2.72,2.82,2.35,PANEL,True)
    text(s,h,x+.22,3.05,2.4,.42,19,c,True,align=PP_ALIGN.CENTER)
    text(s,b,x+.25,3.85,2.32,.8,16,WHITE,False,align=PP_ALIGN.CENTER)
text(s,"Se alguém estiver sendo prejudicado: guarde evidências, denuncie na plataforma e procure um adulto ou autoridade de confiança.",1.0,5.72,11.3,.7,16,MUTED,False,align=PP_ALIGN.CENTER)

#9 campaign
s=base("08 • Campanha da equipe","Nem tudo que você vê é real.",9,"Proposta de campanha para a comunidade escolar")
rect(s,.72,2.3,7.25,3.95,PANEL,True)
text(s,"UM VÍDEO PODE\nPARECER REAL.",1.05,2.8,6.55,1.15,28,WHITE,True,align=PP_ALIGN.CENTER)
text(s,"MAS VOCÊ VERIFICOU?",1.05,4.15,6.55,.55,23,CORAL,True,align=PP_ALIGN.CENTER)
text(s,"PARE → VERIFIQUE → PENSE → COMPARTILHE",1.05,5.32,6.55,.35,15,CYAN,True,align=PP_ALIGN.CENTER)
card(s,8.35,2.3,3.95,1.75,"Formato","Cartaz e publicação para redes da escola.",CYAN,16)
card(s,8.35,4.28,3.95,1.97,"Mensagem","Inteligência artificial exige pensamento crítico.",GREEN,16)

#10 conclusion
s=base("09 • Conclusão","Tecnologia também exige responsabilidade",10,"Deepfakes podem criar conteúdos convincentes, mas nosso comportamento influencia o impacto que eles terão.")
text(s,"“Antes de compartilhar,\nconfira a origem.”",1.0,2.55,7.9,1.55,31,WHITE,True)
rect(s,9.6,2.42,2.1,2.1,PANEL,True)
text(s,"IA\n≠\nVERDADE",9.85,2.82,1.6,1.3,23,CYAN,True,align=PP_ALIGN.CENTER)
text(s,"Proteger dados • verificar informações • respeitar as pessoas • compartilhar com contexto",1.0,5.25,11.2,.55,17,GREEN,True,align=PP_ALIGN.CENTER)

#11 roles and sources
s=base("10 • Equipe e referências","Quem somos e de onde partimos",11,"Projeto da Fase 1 da Maratona Tech 2026")
card(s,.72,2.35,5.7,3.95,"Equipe • 3º ano D","Eduardo Valentin Carneiro\nLuiz Gabriel Alves dos Santos\nThiago Henrique Malta Garcia\nFelype Justino da Silva\nPaulo Ricardo Nunes da Silva\nAdryan Vinicius de Almeida",CYAN,13)
card(s,6.75,2.35,5.65,3.95,"Fontes para aprofundar","ANPD — Radar tecnológico sobre deepfakes\ngov.br/anpd/pt-br/assuntos/noticias\n\nTSE — Deepfake e Eleições 2026\ntse.jus.br/comunicacao/noticias\n\nGoverno Digital — Guia de IA Generativa\ngov.br/governodigital/pt-br/infraestrutura-nacional-de-dados",GREEN,12)
text(s,"Colégio Estadual Marechal Costa e Silva • Cidade Gaúcha, Paraná",.85,6.53,11.5,.35,14,WHITE,True,align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
